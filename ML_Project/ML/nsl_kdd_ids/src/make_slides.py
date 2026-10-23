"""Builds the review slide deck (.pptx) from results/results.json + figures.
Usage: python -m src.make_slides --team "Shweta, Pranav P" --team-no 3"""
import argparse, json, os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

NAVY, GREY = RGBColor(0x1F, 0x3A, 0x5F), RGBColor(0x55, 0x55, 0x55)


def pct(x, d=1): return f"{100 * x:.{d}f}%"


def build(a):
    R = json.load(open(os.path.join(a.results, "results.json")))
    P = Presentation(); P.slide_width, P.slide_height = Inches(13.33), Inches(7.5)
    blank = P.slide_layouts[6]

    def slide(title):
        s = P.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, P.slide_width, Inches(0.95))
        bar.fill.solid(); bar.fill.fore_color.rgb = NAVY; bar.line.fill.background()
        tb = s.shapes.add_textbox(Inches(.5), Inches(.15), Inches(12), Inches(.7)).text_frame
        tb.text = title; tb.paragraphs[0].font.size = Pt(30); tb.paragraphs[0].font.bold = True
        tb.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
        return s

    def bullets(s, items, x=.6, y=1.3, w=12, h=5.6, size=22):
        tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame; tf.word_wrap = True
        for i, t in enumerate(items):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "• " + t; p.font.size = Pt(size); p.space_after = Pt(10); p.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    def pic(s, name, x, y, w):
        s.shapes.add_picture(os.path.join(a.results, name), Inches(x), Inches(y), width=Inches(w))

    # 1 title
    s = P.slides.add_slide(blank)
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, P.slide_width, P.slide_height); bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY; bg.line.fill.background()
    tf = s.shapes.add_textbox(Inches(.8), Inches(2.3), Inches(11.5), Inches(3)).text_frame; tf.word_wrap = True
    tf.text = "Neural Network for Security Attack Detection"; tf.paragraphs[0].font.size = Pt(42)
    tf.paragraphs[0].font.bold = True; tf.paragraphs[0].font.color.rgb = RGBColor(255, 255, 255)
    for t in ["NSL-KDD intrusion detection | UE24CS352A Machine Learning Mini-Project", f"Team {a.team_no}: {a.team}"]:
        p = tf.add_paragraph(); p.text = t; p.font.size = Pt(20); p.font.color.rgb = RGBColor(0xDD, 0xE6, 0xF2)

    s = slide("Problem statement")
    bullets(s, ["Attackers change behaviour, so signature-based intrusion detection misses new attacks",
                "Goal: predict the attack type of a network connection from its recorded parameters",
                "Goal: decide normal vs abnormal behaviour, including attacks never seen in training",
                "Approach: a neural network trained on the NSL-KDD benchmark"])

    s = slide("Dataset: NSL-KDD")
    bullets(s, [f"Train: {R['n_train']:,} records | Test: {R['n_test']:,} records | 41 features each",
                "Feature groups: basic TCP/IP, content, traffic; 3 categorical (protocol, service, flag)",
                "Classes: Normal, DoS, Probe, R2L, U2R (22 attacks in train)",
                f"Test set has {R['n_unseen_test_attack_types']} attack types not in training, which acts as a zero-day test",
                "Improved over KDD Cup 99: no redundant train records, no duplicate test records"], w=6.6, size=19)
    pic(s, "fig_distribution.png", 7.4, 1.6, 5.4)

    s = slide("Approach: preprocessing")
    bullets(s, [f"One-hot encode protocol_type, service, flag: 41 → {R['n_inputs']} input columns",
                "Standardise each column (mean 0, std 1) using training statistics only",
                f"Labels: {R['n_outputs']} outputs = normal + 22 training attacks + 1 'other' unit for unseen attacks",
                "Two metrics: attack-type accuracy, and normal-vs-abnormal accuracy"])

    s = slide("Model architecture")
    xs = [1.2, 5.2, 9.2]
    labels = [(f"Input\n{R['n_inputs']} units", "#5B8DB8"), (f"Hidden\n{R['n_hidden']} units\nsigmoid", "#3C6E9F"),
              (f"Output\n{R['n_outputs']} units\nsoftmax", "#1F3A5F")]
    for x, (t, col) in zip(xs, labels):
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(1.8), Inches(2.6), Inches(1.8))
        b.fill.solid(); b.fill.fore_color.rgb = RGBColor.from_string(col[1:]); b.line.fill.background()
        b.text_frame.text = t
        for p in b.text_frame.paragraphs: p.font.size = Pt(20); p.font.color.rgb = RGBColor(255, 255, 255)
    for x in (3.9, 7.9):
        s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(2.4), Inches(1.2), Inches(.6)).fill.solid()
    bullets(s, [f"Loss: cross-entropy + L2 (λ={R['l2']}) | Optimiser: mini-batch gradient descent",
                f"Learning rate {R['lr']}, batch size {R['batch_size']}, {R['epochs']} epochs",
                "Implemented from scratch in NumPy (forward pass, back-propagation, SGD)"], y=4.3, size=20)

    s = slide("Implementation overview")
    bullets(s, ["src/data.py: load, one-hot, standardise, build labels",
                "src/model.py: NumPy neural network and mini-batch training",
                "src/metrics.py: accuracies and confusion counts",
                "src/train.py: multi-seed runs, epoch curves, baselines, figures, results.json",
                "src/make_report.py / make_slides.py: write-up and deck generated from real results",
                "tests/: unit and smoke tests | run_all.py: one command runs everything"], size=21)

    s = slide("Results: accuracy")
    bullets(s, [f"Attack-type: train {pct(R['train_multi_mean'])}, test {pct(R['test_multi_mean'])} ± {pct(R['test_multi_std'])}",
                f"Normal vs abnormal: train {pct(R['train_bin_mean'])}, test {pct(R['test_bin_mean'])} ± {pct(R['test_bin_std'])}",
                f"Test binary accuracy over {len(R['runs'])} seeds: {pct(R['test_bin_min'])} – {pct(R['test_bin_max'])} (initialisation matters)"]
            + [f"Baseline {k}: {pct(v['test_bin'])} (binary)" for k, v in R.get("baselines", {}).items()], h=2.4, size=20)
    pic(s, "fig_curves.png", 2.3, 4.2, 8.6)

    cm = R["confusion_binary"]
    s = slide("Results: where does it work?")
    pic(s, "fig_confusion.png", .6, 1.4, 5.6); pic(s, "fig_detection.png", 6.8, 1.4, 5.8)
    txt = f"Best seed: false-alarm rate {pct(cm['false_alarm_rate'])}, attack recall {pct(cm['recall'])}, F1 {cm['f1']:.3f}"
    if R.get("unseen_attack_detection") is not None:
        txt += f" | Detected: seen attacks {pct(R['seen_attack_detection'])}, unseen {pct(R['unseen_attack_detection'])}"
    bullets(s, [txt], y=5.8, h=1.4, size=18)

    s = slide("Conclusions")
    bullets(s, ["A 122-500-24 network learns the training data well; test accuracy is lower because of unseen attacks",
                "Framing the task as normal vs abnormal gives a more useful detection number than exact attack type",
                "Results vary with random initialisation, so we report several seeds",
                "Limitations: rare U2R/R2L classes, distribution shift, false alarms need human review",
                "Future work: class weighting, early stopping, more datasets"], size=21)

    s = slide("Live demo")
    bullets(s, ["1. pip install -r requirements.txt", "2. Place KDDTrain+.txt and KDDTest+.txt in data/",
                "3. python run_all.py --team \"Shweta (PES1UG24AM404), Pranav P (PES1UG24AM387)\" --team-no 3", "4. Outputs: results/results.json, figures, docs/*.pdf|pptx"], size=22)
    os.makedirs(os.path.dirname(a.out) or ".", exist_ok=True); P.save(a.out); print("wrote", a.out)


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--results", default="results"); p.add_argument("--out", default="docs/Project_Slides.pptx")
    p.add_argument("--team", default="Shweta (PES1UG24AM404), Pranav P (PES1UG24AM387)"); p.add_argument("--team-no", default="3")
    build(p.parse_args())
